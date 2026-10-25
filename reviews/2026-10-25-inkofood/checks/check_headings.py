#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сверка: заголовки разделов пояснения совпадают с заголовками слайдов и идут в том же порядке."""
import re, sys
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
prs = Presentation(HERE / "Inkofood_tri_agent_synthesis_2026.pptx")
md = (HERE / "Inkofood_tri_agent_synthesis_2026_poyasnenie.md").read_text(encoding="utf-8")
heads = re.findall(r"^## Слайд (\d+)\. (.+)$", md, re.M)
bad = []
for i, (n, ht) in enumerate(heads, 1):
    ts = [sh.text_frame.text.strip() for sh in prs.slides[i - 1].shapes
          if sh.has_text_frame and sh.text_frame.text.strip()]
    title = re.sub(r"\s+", " ", ts[1] if len(ts) > 1 else ts[0]).strip()
    h = re.sub(r"\s+", " ", ht).strip()
    if int(n) != i or title.lower() not in h.lower():
        bad.append((i, n, title[:60], h[:60]))
print(f"слайдов: {len(prs.slides)}, разделов: {len(heads)}, несовпадений: {len(bad)}")
for b in bad:
    print("  ", b)
sys.exit(1 if bad else 0)
